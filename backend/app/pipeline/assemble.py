"""Stage 6: assemble the Word document from template pack + case record + approved statements.

Review copy:
  - approved clauses sit in *locked* content controls tagged ``clause:<id>@<version>``;
  - AI-drafted text sits in content controls tagged ``ai:<slot>:<statement ids>`` in the "AI Draft"
    style, with statement reference markers that match the evidence appendix;
  - INSUFFICIENT_EVIDENCE markers are shown in red.
Client copy: same content, house styles only, no markers. Refuses to render if anything is missing.
Evidence appendix: internal record of statement -> fact -> source for the file.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Literal

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.domain.catalogue import Catalogue
from app.domain.schemas import CaseRecord, Document, Report, SourceSpan, Statement
from app.pipeline.draft import section_included
from app.templates_mgmt.pack import PLACEHOLDER, AISlotBlock, ClauseBlock, FactsTableBlock, TemplatePack

Mode = Literal["review", "client"]
AI_STYLE, CLAUSE_STYLE = "AI Draft", "Approved Wording"


class AssemblyError(RuntimeError):
    pass


def _ensure_styles(d: docx.document.Document) -> None:
    names = {s.name for s in d.styles}
    if AI_STYLE not in names:
        st = d.styles.add_style(AI_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = d.styles["Normal"]
        st.font.color.rgb = RGBColor(0x5B, 0x2C, 0x83)
    if CLAUSE_STYLE not in names:
        st = d.styles.add_style(CLAUSE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = d.styles["Normal"]
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def _new_document() -> docx.document.Document:
    """Blank document in UK A4 with 2.5 cm margins (python-docx defaults to US Letter)."""
    d = docx.Document()
    for section in d.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)
    return d


def _wrap_in_content_control(paragraph: Paragraph, *, tag: str, alias: str, locked: bool) -> None:
    """Move a paragraph inside a block-level structured document tag (Word content control)."""
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    a = OxmlElement("w:alias")
    a.set(qn("w:val"), alias)
    t = OxmlElement("w:tag")
    t.set(qn("w:val"), tag[:64])  # Word limits tag length
    pr.extend([a, t])
    if locked:
        lock = OxmlElement("w:lock")
        lock.set(qn("w:val"), "sdtContentLocked")
        pr.append(lock)
    content = OxmlElement("w:sdtContent")
    sdt.extend([pr, content])
    p = paragraph._p
    p.addprevious(sdt)
    content.append(p)


def _render_text(text: str, record: CaseRecord, catalogue: Catalogue, mode: Mode) -> str:
    def sub(m):
        key = m[1]
        fact = record.usable_fact(key)
        if fact is None:
            if mode == "client":
                raise AssemblyError(f"no verified value for {key}")
            return f"[MISSING: {catalogue.definition(key).label}]"
        return fact.display

    return PLACEHOLDER.sub(sub, text)


def render_report(
    pack: TemplatePack,
    record: CaseRecord,
    report: Report,
    catalogue: Catalogue,
    *,
    mode: Mode,
    layout: Path | None = None,
) -> bytes:
    if mode == "client" and any(s.insufficient_evidence for s in report.statements):
        raise AssemblyError("cannot produce client copy while INSUFFICIENT_EVIDENCE markers remain")

    d = docx.Document(str(layout)) if layout and layout.exists() else _new_document()
    _ensure_styles(d)
    d.core_properties.title = pack.name
    d.core_properties.comments = (
        f"report={report.id} template={pack.id}@{pack.version} sha256={pack.content_hash[:16]} "
        f"model={report.drafting_model} prompt={report.prompt_version} revision={report.revision}"
    )
    if mode == "review":
        header = d.sections[0].header.paragraphs[0]
        header.text = f"DRAFT FOR INTERNAL REVIEW: {report.id} rev {report.revision}. Not for issue."

    d.add_heading(pack.name, level=0)
    if mode == "review":
        key = d.add_paragraph()
        key.add_run("Key: ").bold = True
        key.add_run("approved wording (locked) ").font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        key.add_run("· AI-drafted text, [Sn] = evidence appendix reference ").font.color.rgb = RGBColor(
            0x5B, 0x2C, 0x83)
        key.add_run("· INSUFFICIENT EVIDENCE blocks issue").font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    by_slot: dict[str, list[tuple[int, Statement]]] = {}
    for n, s in enumerate(report.statements, start=1):
        by_slot.setdefault(s.slot_id, []).append((n, s))

    for section in pack.sections:
        if not section_included(section, record):
            continue
        d.add_heading(section.heading, level=1)
        for block in section.blocks:
            if isinstance(block, ClauseBlock):
                clause = pack.clauses[block.clause]
                text = _render_text(clause.text, record, catalogue, mode)
                for chunk in [c for c in text.split("\n\n") if c.strip()]:
                    p = d.add_paragraph(chunk.strip(), style=CLAUSE_STYLE if mode == "review" else None)
                    if mode == "review":
                        _wrap_in_content_control(p, tag=f"clause:{clause.id}@{clause.version}",
                                                 alias=f"Approved wording {clause.id} v{clause.version}",
                                                 locked=True)
            elif isinstance(block, FactsTableBlock):
                table = d.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                for key in block.rows:
                    fact = record.usable_fact(key)
                    if fact is None and mode == "client":
                        raise AssemblyError(f"no verified value for {key}")
                    row = table.add_row().cells
                    row[0].text = catalogue.definition(key).label
                    row[1].text = fact.display if fact else "[MISSING]"
                d.add_paragraph()
            elif isinstance(block, AISlotBlock):
                items = by_slot.get(block.slot, [])
                if not items:
                    continue
                p = d.add_paragraph(style=AI_STYLE if mode == "review" else None)
                for n, s in items:
                    run = p.add_run(s.text.strip() + " ")
                    if s.insufficient_evidence:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                    if mode == "review":
                        ref = p.add_run(f"[S{n}] ")
                        ref.font.superscript = True
                        ref.font.size = Pt(7)
                if mode == "review":
                    ids = ",".join(s.id.removeprefix("stmt_")[:6] for _, s in items)
                    _wrap_in_content_control(p, tag=f"ai:{block.slot}:{ids}",
                                             alias=f"AI draft: {block.slot}", locked=False)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_evidence_appendix(
    report: Report,
    record: CaseRecord,
    catalogue: Catalogue,
    spans: dict[str, SourceSpan],
    documents: dict[str, Document],
) -> bytes:
    d = _new_document()
    d.add_heading("Evidence appendix (internal, not for issue)", level=0)
    d.add_paragraph(
        f"Report {report.id}, revision {report.revision}, template {report.template_pack_id}@"
        f"{report.template_version} ({report.template_hash[:16]}), drafting model {report.drafting_model}, "
        f"prompt {report.prompt_version}, case record revision {report.case_record_revision}."
    )
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["Ref", "Statement", "Fact (verified by)", "Source"], strict=True):
        cell.text = title
    for n, s in enumerate(report.statements, start=1):
        facts = [f for fid in s.fact_ids if (f := record.fact(fid))]
        if not facts:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text, row[3].text = f"S{n}", s.text, "none", "none"
            continue
        for i, fact in enumerate(facts):
            sources = []
            for src in fact.sources:
                if src.kind == "span" and (span := spans.get(src.span_id or "")):
                    doc = documents[span.document_id]
                    sources.append(f'{doc.filename} (sha256 {doc.sha256[:12]}…) p{span.page}: "{span.text}"')
                else:
                    sources.append(f"Adviser input by {src.recorded_by}" + (f": {src.note}" if src.note else ""))
            row = table.add_row().cells
            row[0].text = f"S{n}" if i == 0 else ""
            row[1].text = s.text if i == 0 else ""
            row[2].text = f"{catalogue.definition(fact.key).label} = {fact.display} ({fact.verified_by})"
            row[3].text = "\n".join(sources)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
