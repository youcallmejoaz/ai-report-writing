import io

import docx
import pytest

from app.pipeline.assemble import AssemblyError, render_report
from tests.conftest import prepare_ready_case

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _controls(data):
    body = docx.Document(io.BytesIO(data)).element.body
    out = []
    for sdt in body.iter(f"{W}sdt"):
        tag = sdt.find(f".//{W}tag").get(f"{W}val")
        locked = sdt.find(f".//{W}lock") is not None
        out.append((tag, locked))
    return out


def _text(data):
    return "\n".join(t.text or "" for t in docx.Document(io.BytesIO(data)).element.body.iter(f"{W}t"))


@pytest.fixture
def report(service, users, pack):
    case = prepare_ready_case(service, users, pack.id)
    return service.generate_report(users["alice"], case.id)


def test_review_copy_marks_provenance(service, users, report):
    controls = _controls(service.render(users["alice"], report.id, "review"))
    clauses = [c for c in controls if c[0].startswith("clause:")]
    ai = [c for c in controls if c[0].startswith("ai:")]
    assert clauses and all(locked for _, locked in clauses)
    assert {t.split(":")[1] for t, _ in ai} == report.ai_slots
    assert not any(locked for _, locked in ai)
    assert ("clause:risk_warnings@2.0.0", True) in controls


def test_client_copy_is_clean(service, users, report):
    data = service.render(users["alice"], report.id, "client")
    text = _text(data)
    assert _controls(data) == []
    assert "[S1]" not in text and "INTERNAL REVIEW" not in text
    assert "Dear Jane Example" in text and "£125,400.00" in text
    assert "Safeguarded benefits" not in text.split("Your existing pension")[1]  # conditional section omitted


def test_client_copy_refuses_insufficient_evidence(service, users, pack, report, catalogue):
    report.statements[0].insufficient_evidence = True
    record = service.store.records[report.case_id]
    with pytest.raises(AssemblyError):
        render_report(pack, record, report, catalogue, mode="client")


def test_evidence_appendix_links_statements_to_sources(service, users, report):
    text = _text(service.render(users["alice"], report.id, "evidence"))
    assert "pension_statement.txt" in text and "Total fund value: £125,400.00" in text
    assert "Adviser input by aaron" in text


def test_documents_are_a4(service, users, report):
    for variant in ("review", "client", "evidence"):
        section = docx.Document(io.BytesIO(service.render(users["alice"], report.id, variant))).sections[0]
        assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
